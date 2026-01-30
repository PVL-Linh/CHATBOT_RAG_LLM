import os
import fitz # PyMuPDF
import re
from typing import Dict, List
# Import hàm format mới từ file Data_processing
from Data_processing import clean_page_text 

# Config
USE_OCR_FALLBACK = True
MIN_OCR_LEN = 40

def _process_pdf(pdf_path: str) -> str:
    """Xử lý PDF: Extract text -> Format Markdown"""
    doc = fitz.open(pdf_path)
    page_texts = []
    
    for page in doc:
        raw = page.get_text()
        
        # Logic OCR fallback (giản lược)
        if len(raw.strip()) < MIN_OCR_LEN and USE_OCR_FALLBACK:
             # (Tại đây bạn có thể gọi code OCR nếu muốn)
             pass
             
        # Gọi hàm clean đã update logic Markdown
        cleaned = clean_page_text(raw, drop_headers=True)
        page_texts.append(cleaned)
        
    doc.close()
    return "\n\n".join(page_texts)

def _process_docx(docx_path: str) -> str:
    """Xử lý DOCX: Extract text -> Format Markdown"""
    try:
        import docx
        doc = docx.Document(docx_path)
        full_text = []
        for para in doc.paragraphs:
            clean = clean_page_text(para.text, drop_headers=False)
            full_text.append(clean)
        return "\n".join(full_text)
    except Exception as e:
        print(f"Lỗi đọc docx {docx_path}: {e}")
        return ""

def process_documents(input_dir: str, output_dir: str):
    """
    Quét thư mục input, convert PDF/DOCX -> .md
    """
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    for root, _, files in os.walk(input_dir):
        # Tạo cấu trúc thư mục tương ứng ở output
        rel_path = os.path.relpath(root, input_dir)
        target_dir = os.path.join(output_dir, rel_path)
        os.makedirs(target_dir, exist_ok=True)

        for filename in files:
            file_path = os.path.join(root, filename)
            
            # Đổi đuôi file thành .md
            name_no_ext = os.path.splitext(filename)[0]
            md_filename = f"{name_no_ext}.md"
            output_path = os.path.join(target_dir, md_filename)

            content = ""
            if filename.lower().endswith(".pdf"):
                print(f"📄 Processing PDF: {filename} -> .md")
                content = _process_pdf(file_path)
            elif filename.lower().endswith(".docx"):
                print(f"📄 Processing DOCX: {filename} -> .md")
                content = _process_docx(file_path)
            
            if content:
                with open(output_path, "w", encoding="utf-8") as f:
                    f.write(content)

# Wrapper function để gọi từ main
def processing_Data_doclinkToText(input_dir, output_dir, **kwargs):
    print("🚀 Bắt đầu chuyển đổi tài liệu sang Markdown...")
    process_documents(input_dir, output_dir)
    print("✅ Hoàn tất chuyển đổi.")



# import os
# import zipfile
# from typing import Dict, List, Tuple
# import fitz
# import re
# from Data_processing import clean_page_text
# try:
#     from lxml import etree
#     LXML_OK = True
# except Exception:
#     LXML_OK = False

# import xml.etree.ElementTree as ET
# from .batch_tree_v2_local import tree_run
# from .docx_org_tree_local import org_tree_run
# try:
#     import docx as _docx
# except Exception:
#     raise ImportError("Chưa cài python-docx. Cài: pip install python-docx lxml")

# USE_OCR_FALLBACK = True
# DOCX_DROP_HEADERS = False
# PDF_DROP_HEADERS = True
# TABLE_MODE = "flatten"
# KEEP_TABLES = True
# LATEX_MATH = True
# DEDUP_NORMALIZE = True
# DEDUP_LOWERCASE = True
# MIN_OCR_LEN = int(os.getenv("MIN_OCR_LEN", 40))
# OCR_DPI = int(os.getenv("OCR_DPI", 200))

# try:
#     import easyocr
#     import numpy as np
#     from PIL import Image
#     import io
#     _easyocr_reader = easyocr.Reader(["vi", "en"])
# except Exception:
#     USE_OCR_FALLBACK = False
#     _easyocr_reader = None

# def _normalize_line(s: str) -> str:
#     if s is None:
#         return ""
#     x = s
#     x = x.replace("\u00A0", " ")
#     x = x.replace("–", "-").replace("—", "-").replace("’", "'").replace("“", '"').replace("”", '"')
#     x = re.sub(r"[ \t]+", " ", x)
#     x = re.sub(r"\s+\n", "\n", x)
#     x = re.sub(r"\n\s+", "\n", x)
#     x = x.strip()
#     if DEDUP_LOWERCASE:
#         x = x.lower()
#     return x

# def _uniq_keep_order_norm(lines: List[str]) -> List[str]:
#     if not DEDUP_NORMALIZE:
#         seen = set()
#         out = []
#         for s in lines:
#             if s not in seen:
#                 seen.add(s)
#                 out.append(s)
#         return out

#     seen_norm = set()
#     out = []
#     for s in lines:
#         key = _normalize_line(s)
#         if key and key not in seen_norm:
#             seen_norm.add(key)
#             out.append(s.strip())
#     return out

# def _extract_text(page: fitz.Page) -> str:
#     try:
#         data = page.get_text("dict")
#         out = []
#         for b in data.get("blocks", []):
#             for ln in b.get("lines", []):
#                 out.append("".join(sp.get("text", "") for sp in ln.get("spans", [])))
#         txt = "\n".join(out)
#         if txt.strip():
#             return txt
#     except Exception:
#         pass
#     return page.get_text() or ""

# def _ocr_image(pix: fitz.Pixmap) -> str:
#     """OCR 1 trang PDF bằng EasyOCR."""
#     if not USE_OCR_FALLBACK or _easyocr_reader is None:
#         return ""
#     img_bytes = pix.tobytes("png")
#     img = np.array(Image.open(io.BytesIO(img_bytes)))
#     try:
#         results = _easyocr_reader.readtext(img, detail=0, paragraph=True)
#     except Exception:
#         results = []
#     return "\n".join(results)


# def _process_pdf(pdf_path: str) -> str:
#     doc = fitz.open(pdf_path)
#     page_texts: List[str] = []

#     for page in doc:
#         raw = _extract_text(page)

#         if USE_OCR_FALLBACK and len(raw.strip()) < MIN_OCR_LEN:
#             try:
#                 raw = (raw + "\n" + _ocr_image(page.get_pixmap(dpi=200))).strip()
#             except Exception as e:
#                 print(f"⚠️ OCR lỗi tại {pdf_path}: {e}")

#         cleaned = clean_page_text(
#             raw,
#             drop_headers=PDF_DROP_HEADERS,
#             keep_tables=KEEP_TABLES,
#             table_mode=TABLE_MODE,
#             latex_math=LATEX_MATH,
#         )
#         page_texts.append(cleaned)

#     doc.close()
#     combined = "\n\n".join([p for p in page_texts if p.strip()])

#     if not combined.strip() and USE_OCR_FALLBACK:
#         print(f"⚠️ File PDF rỗng, chạy OCR toàn file: {pdf_path}")
#         all_text = []
#         for page in fitz.open(pdf_path):
#             ocr = _ocr_image(page.get_pixmap(dpi=200))
#             if ocr.strip():
#                 all_text.append(ocr)
#         combined = "\n\n".join(all_text)

#     return combined

# def _table_lines(tbl) -> List[str]:
#     lines = []
#     for row in tbl.rows:
#         txt = "\t".join((cell.text or "").strip() for cell in row.cells).strip()
#         if txt:
#             lines.append(txt)
#     return lines


# def _hf_to_lines(hf) -> List[str]:
#     if hf is None:
#         return []
#     lines = []
#     for p in hf.paragraphs:
#         t = (p.text or "").strip()
#         if t:
#             lines.append(t)
#     for tbl in hf.tables:
#         lines.extend(_table_lines(tbl))
#     return lines


# def _extract_hf_python_docx(document) -> Tuple[List[str], List[str]]:
#     headers, footers = [], []
#     for sec in document.sections:
#         headers += _hf_to_lines(sec.header)
#         footers += _hf_to_lines(sec.footer)
#         try:
#             headers += _hf_to_lines(sec.first_page_header)
#             footers += _hf_to_lines(sec.first_page_footer)
#         except Exception:
#             pass
#         try:
#             headers += _hf_to_lines(sec.even_page_header)
#             footers += _hf_to_lines(sec.even_page_footer)
#         except Exception:
#             pass
#     return headers, footers


# def _uniq_keep_order(lines: List[str]) -> List[str]:
#     seen = set()
#     out = []
#     for s in lines:
#         if s not in seen:
#             seen.add(s)
#             out.append(s)
#     return out


# def _extract_hf_xml(docx_path: str) -> Tuple[List[str], List[str]]:
#     if not LXML_OK:
#         return [], []
#     ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
#     headers, footers = [], []
#     with zipfile.ZipFile(docx_path) as z:
#         for name in z.namelist():
#             if name.startswith("word/header") and name.endswith(".xml"):
#                 root = etree.fromstring(z.read(name))
#                 for p in root.xpath(".//w:p", namespaces=ns):
#                     txt = "".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
#                     if txt:
#                         headers.append(txt)
#             if name.startswith("word/footer") and name.endswith(".xml"):
#                 root = etree.fromstring(z.read(name))
#                 for p in root.xpath(".//w:p", namespaces=ns):
#                     txt = "".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
#                     if txt:
#                         footers.append(txt)
#     return headers, footers


# def _extract_body_textboxes_xml(docx_path: str) -> List[str]:
#     if not LXML_OK:
#         return []
#     ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
#     lines = []
#     with zipfile.ZipFile(docx_path) as z:
#         if "word/document.xml" not in z.namelist():
#             return []
#         root = etree.fromstring(z.read("word/document.xml"))
#         # CHỈ bắt p bên trong textbox content
#         for p in root.xpath(".//w:txbxContent//w:p", namespaces=ns):
#             txt = "".join(p.xpath(".//w:t/text()", namespaces=ns)).strip()
#             if txt:
#                 lines.append(txt)
#     return lines


# def _extract_top_tables_as_rows(docx_path: str, max_tables: int = 2) -> List[str]:
#     try:
#         ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
#         with zipfile.ZipFile(docx_path) as zf:
#             root = ET.fromstring(zf.read('word/document.xml'))
#         tables = root.findall('.//w:tbl', ns)[:max_tables]

#         def norm(s: str) -> str:
#             s = s or ''
#             s = s.replace("\u00A0", " ")
#             s = ' '.join(s.split())
#             return s

#         lines = []
#         for tbl in tables:
#             for tr in tbl.findall('.//w:tr', ns):
#                 cells = []
#                 for tc in tr.findall('.//w:tc', ns):
#                     texts = [t.text for t in tc.findall('.//w:t', ns) if t.text]
#                     cells.append(norm(''.join(texts)))
#                 line = ' \t '.join(c for c in cells if c)
#                 if line.strip():
#                     lines.append(line)
#         return lines
#     except Exception:
#         return []


# def _process_docx(docx_path: str) -> str:
#     """Xử lý file DOCX → text (paragraph + bảng + header/footer, có XML fallback + head từ bảng)."""
#     document = _docx.Document(docx_path)
#     hdr_lines, ftr_lines = _extract_hf_python_docx(document)
#     if LXML_OK and (not hdr_lines or not ftr_lines):
#         hdr2, ftr2 = _extract_hf_xml(docx_path)
#         if not hdr_lines:
#             hdr_lines = hdr2
#         if not ftr_lines:
#             ftr_lines = ftr2
#     top_table_head = _extract_top_tables_as_rows(docx_path, max_tables=1)
#     para_lines = [(p.text or "").strip() for p in document.paragraphs if (p.text or "").strip()]
#     body_tables = []
#     for tbl in document.tables:
#         body_tables.extend(_table_lines(tbl))

#     body_from_xml = _extract_body_textboxes_xml(docx_path)

#     parts = []
#     if top_table_head:
#         parts.extend(top_table_head)
#     if not DOCX_DROP_HEADERS and hdr_lines:
#         parts.extend(hdr_lines)
#     parts.extend(para_lines)
#     if body_tables:
#         parts.extend(body_tables)
#     if body_from_xml:
#         parts.extend(body_from_xml)
#     if not DOCX_DROP_HEADERS and ftr_lines:
#         parts.extend(ftr_lines)

#     all_lines = _uniq_keep_order_norm([ln for ln in parts if ln and ln.strip()])
#     combined = "\n".join(all_lines)

#     return clean_page_text(
#         combined,
#         drop_headers=DOCX_DROP_HEADERS,
#         keep_tables=KEEP_TABLES,
#         table_mode=TABLE_MODE,
#         latex_math=LATEX_MATH,
#     )

# def process_documents(input_folder: str, output_folder: str = None) -> Dict[str, str]:
#     """
#     Trích xuất toàn bộ PDF/DOCX → TXT (duyệt cả thư mục con).
#     Nếu output_folder=None thì file TXT sẽ được lưu cùng thư mục với file gốc.
#     Nếu có output_folder thì giữ nguyên cấu trúc thư mục con bên trong.
#     Trả về dict {txt_relpath: content}.
#     """
#     input_folder = os.path.abspath(input_folder)
#     if output_folder:
#         output_folder = os.path.abspath(output_folder)

#     results: Dict[str, str] = {}

#     for root, _, files in os.walk(input_folder):
#         for filename in files:
#             if filename.startswith("~$"):
#                 continue
#             if not filename.lower().endswith((".pdf", ".docx")):
#                 continue

#             file_path = os.path.join(root, filename)
#             rel_path = os.path.relpath(root, input_folder)
#             txt_filename = os.path.splitext(filename)[0] + ".txt"

#             if output_folder:
#                 target_dir = os.path.join(output_folder, rel_path)
#             else:
#                 target_dir = root

#             os.makedirs(target_dir, exist_ok=True)
#             txt_path = os.path.join(target_dir, txt_filename)

#             print(f"🔍 Đang xử lý: {file_path}")

#             if filename.lower().endswith(".pdf"):
#                 combined = _process_pdf(file_path)
#             else:
#                 combined = _process_docx(file_path)

#             with open(txt_path, "w", encoding="utf-8") as f:
#                 f.write(combined)

#             results[os.path.join(rel_path, txt_filename)] = combined

#     print("✅ Hoàn tất xử lý tất cả file PDF/DOCX.")
#     return results

# def processing_Data_doclinkToText(input_dir = "./Documents/HR/txt", output_dir = "./src/app/Data/HR/txt", 
#                                 input_txt_diagram = "./Documents/HR/Diagram/documents_workflowAndText", 
#                                 output_txt_diagram = "./src/app/Data/HR/Diagram/documents_workflowAndText",
#                                 input_dir_org_tree = "./Documents/HR/Diagram/procedure",
#                                 output_dir_org_tree = "./src/app/Data/HR/Diagram/procedure"):
#     process_documents(input_dir, output_dir)
#     stats = tree_run(input_txt_diagram, output_txt_diagram)
#     org_tree_run(input_dir_org_tree, output_dir_org_tree)

# def processing_Data_dockinkToText_v2 (input_dir, output_dir):
#     process_documents(input_dir, output_dir)

