import os
from openpyxl import load_workbook
from docx import Document
import fitz  # PyMuPDF


# ================== XỬ LÝ XLSX -> TXT THEO CỘT ==================
def xlsx_to_txt_by_column_all_sheets(input_path: str, output_dir: str):
    """
    Convert 1 file .xlsx thành nhiều file .txt (mỗi sheet 1 file),
    dữ liệu ghi THEO CỘT.
    - Hàng hoàn toàn trống: bỏ
    - Ô trống trong hàng còn dữ liệu: thay bằng "-"
    - Header của cột = ô đầu tiên khác "-" / rỗng
    """
    os.makedirs(output_dir, exist_ok=True)

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    wb = load_workbook(input_path, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Làm sạch tên sheet để dùng trong tên file
        safe_sheet_name = "".join(
            c if c.isalnum() or c in (" ", "_", "-") else "_"
            for c in sheet_name
        ).strip()

        txt_filename = f"{base_name}__{safe_sheet_name}.txt"
        txt_path = os.path.join(output_dir, txt_filename)

        rows = []
        for row in ws.iter_rows(values_only=True):
            raw_values = ["" if v is None else str(v).strip() for v in row]

            # Bỏ hẳn hàng trống
            if all(v == "" for v in raw_values):
                continue

            # Ô trống -> "-"
            values = [v if v != "" else "-" for v in raw_values]
            rows.append(values)

        if not rows:
            continue

        col_count = max(len(r) for r in rows)

        with open(txt_path, "w", encoding="utf-8-sig") as f_txt:
            for col_idx in range(col_count):
                col_values = []
                for r in rows:
                    value = r[col_idx] if col_idx < len(r) else "-"
                    if value == "":
                        value = "-"
                    col_values.append(value)

                # Nếu cả cột toàn "-" thì bỏ
                if all(v == "-" for v in col_values):
                    continue

                # ====== TÌM HEADER: ô đầu tiên khác "-" / rỗng ======
                header = None
                header_idx = None
                for i, v in enumerate(col_values):
                    if v not in ("", "-"):
                        header = v
                        header_idx = i
                        break
                if header is None:
                    header = f"Column_{col_idx + 1}"

                f_txt.write(f"[{header}]\n")

                # Ghi các giá trị còn lại, bỏ dòng header ra
                for i, v in enumerate(col_values):
                    if i == header_idx:
                        continue
                    f_txt.write(v + "\n")

                f_txt.write("\n")

        print(f"✔ XLSX -> TXT theo cột (header thông minh): {txt_path}")

# ================== XỬ LÝ PDF -> TXT ==================
def pdf_to_txt(input_path: str, output_dir: str):
    """
    Convert 1 file PDF thành 1 file .txt (text thường).
    """
    os.makedirs(output_dir, exist_ok=True)

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    txt_path = os.path.join(output_dir, base_name + ".txt")

    doc = fitz.open(input_path)
    texts = []
    for page in doc:
        texts.append(page.get_text())
    doc.close()

    with open(txt_path, "w", encoding="utf-8-sig") as f:
        f.write("\n\n".join(texts))

    print(f"✔ PDF -> TXT: {txt_path}")


# ================== XỬ LÝ DOCX -> TXT ==================
def docx_to_txt(input_path: str, output_dir: str):
    """
    Convert 1 file .docx thành 1 file .txt (text thường).
    """
    os.makedirs(output_dir, exist_ok=True)

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    txt_path = os.path.join(output_dir, base_name + ".txt")

    doc = Document(input_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip() != ""]

    with open(txt_path, "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines))

    print(f"✔ DOCX -> TXT: {txt_path}")


# ================== BATCH TOÀN BỘ THƯ MỤC (ĐỆ QUY) ==================
def batch_convert_to_txt(input_dir: str, output_dir: str):
    """
    Duyệt toàn bộ cây thư mục input_dir (đệ quy):
    - .xlsx  -> TXT theo cột (mỗi sheet 1 file)
    - .pdf   -> TXT thường (1 file)
    - .docx  -> TXT thường (1 file)
    - Bỏ qua file tạm ~$.xlsx
    Cấu trúc thư mục trong output_dir giống y input_dir (Team Hàn, Team Nhật,...).
    """
    for root, dirs, files in os.walk(input_dir):
        for filename in files:
            if filename.startswith("~$"):  # file lock của Excel
                print(f"Bỏ qua file tạm: {filename}")
                continue

            ext = os.path.splitext(filename)[1].lower()
            src_path = os.path.join(root, filename)

            # Tính thư mục output tương ứng (giữ cấu trúc thư mục)
            rel_dir = os.path.relpath(root, input_dir)  # ví dụ: "Team Nhật"
            out_dir = os.path.join(output_dir, rel_dir)
            os.makedirs(out_dir, exist_ok=True)

            if ext == ".xlsx":
                print(f"Đang xử lý XLSX: {src_path}")
                xlsx_to_txt_by_column_all_sheets(src_path, out_dir)

            elif ext == ".pdf":
                print(f"Đang xử lý PDF:  {src_path}")
                pdf_to_txt(src_path, out_dir)

            elif ext == ".docx":
                print(f"Đang xử lý DOCX: {src_path}")
                docx_to_txt(src_path, out_dir)

            else:
                print(f"Bỏ qua file không hỗ trợ: {src_path}")


if __name__ == "__main__":
    # Đường dẫn tới thư mục Sales của bạn
    # Ví dụ: r"D:\Project\CHATBOT_TIXIMAX_ALL\src\app\Data\Documents\Sales"
    INPUT_DIR = r"Documents\Sales"

    # Thư mục để chứa toàn bộ txt đã convert
    OUTPUT_DIR = r"src\app\Data\Sales"

    batch_convert_to_txt(INPUT_DIR, OUTPUT_DIR)
    print("Hoàn tất.")
