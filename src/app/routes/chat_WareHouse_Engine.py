from __future__ import annotations
import os
import re
import time
import json
from typing import List, Dict, Any, Tuple, Optional
import fitz
import docx as docx_lib
import unicodedata
from pathlib import Path
from langchain_core.documents import Document
from langchain_community.vectorstores import FAISS
from flask import session
from app.services.history import get_history, add_message, create_new_session
from app.Helpers.prompt_internal import SYSTEM_PRIMER_WAREHOUSE, sys_instr_warehouse
from app.Model_LLM.model_llm import LLM_model
from app.Model_LLM.hybrid_retriever import rerank, TOP_K
from app.config.settings import ChatConfig, Chat_engine
from app.Model_LLM.Chat_Database.db_router import handle_db_message
from app.Model_LLM.Chat_Database.redis_ctx import get_ctx, update_ctx, redis_client
from app.config.paths import DATA_DIR_WAREHOUSE, FAISS_ALL_DIR_WAREHOUSE

LLM_SEM = ChatConfig.LLM_SEM
DB_CONF_THRESHOLD = Chat_engine.DB_CONF_THRESHOLD
REWRITE_DB_WITH_LLM = Chat_engine.REWRITE_DB_WITH_LLM
SESSION_DOC_VS: Dict[str, FAISS] = {}


# ======================================================================
# 1. HISTORY / CTX HELPERS
# ======================================================================
def _normalize_vi(text: str) -> str:
    if not text:
        return ""
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text


def _get_history_msgs_for_ctx() -> List[Dict[str, str]]:
    try:
        hist = [
            {"role": m["role"], "content": m["content"]}
            for m in get_history()
            if m["role"] in ("user", "assistant")
        ]
    except Exception:
        hist = []
    return hist


def _get_last_assistant_message(history_msgs: List[Dict[str, str]]) -> Optional[str]:
    for m in reversed(history_msgs or []):
        if m.get("role") == "assistant":
            content = (m.get("content") or "").strip()
            if content:
                return content
    return None


def _load_ctx(session_id: str) -> Dict[str, Any]:
    try:
        ctx = get_ctx(session_id) or {}
        if not isinstance(ctx, dict):
            return {}
        return ctx
    except Exception:
        return {}


# ======================================================================
# 2. DETECT NGÔN NGỮ ĐƠN GIẢN
# ======================================================================
def _auto_detect_lang(text: str) -> str:
    text = text or ""

    vi_chars = re.findall(
        r"[àáạảãăằắặẳẵâầấậẩẫèéẹẻẽêềếệểễ"
        r"ìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũư"
        r"ừứựửữỳýỵỷỹđÀÁẠẢÃĂẰẮẶẲẴÂẦẤẬẨẪ"
        r"ÈÉẸẺẼÊỀẾỆỂỄÌÍỊỈĨÒÓỌỎÕÔỒỐỘỔỖƠ"
        r"ỜỚỢỞỠÙÚỤỦŨƯỪỨỰỬỮỲÝỴỶỸĐ]",
        text,
    )
    if len(vi_chars) >= 1:
        return "vi"

    if re.search(r"[A-Za-z]", text):
        return "en"

    return "vi"


# ======================================================================
# 3. EDITOR: DỊCH/VIẾT LẠI CÂU TRẢ LỜI TRƯỚC
# ======================================================================
def _contents_to_ollama_messages(contents):
    msgs = []
    for m in contents or []:
        role = m.get("role", "user")
        if role not in ("user", "assistant", "system"):
            role = "user"

        text_parts = []
        for p in m.get("parts") or []:
            if isinstance(p, dict) and p.get("text"):
                text_parts.append(p["text"])
        if not text_parts:
            continue

        msgs.append({"role": role, "content": "\n".join(text_parts)})
    return msgs


def _estimate_from_contents(contents, max_out_tokens=1024) -> int:
    words = 0
    for m in contents or []:
        for p in (m.get("parts") or []):
            if isinstance(p, dict) and p.get("text"):
                words += len(p["text"].split())
    return int(1.3 * words) + int(max_out_tokens or 512)


def _safe_gemini_generate(gclient, model, contents, config, retries=3, backoff=0.4):
    from app.Helpers.rate_limit import get_text_limiter

    limiter = get_text_limiter()
    max_out = (
        config.get("max_output_tokens")
        if isinstance(config, dict)
        else getattr(config, "max_output_tokens", 1024)
    )
    tokens_est = _estimate_from_contents(contents, max_out)
    last_err: Optional[Exception] = None

    for i in range(retries + 1):
        try:
            limiter.acquire(tokens_est)
            t0 = time.time()
            with LLM_SEM:
                resp = gclient.models.generate_content(
                    model=model, contents=contents, config=config
                )
            limiter.on_success()
            return (getattr(resp, "text", "") or ""), time.time() - t0
        except Exception as e:
            last_err = e
            s = str(e).lower()
            if ("429" in s or "quota" in s or "resource_exhausted" in s) and i < retries:
                limiter.on_429()
                time.sleep(backoff * (2**i))
                continue
            limiter.on_429()
            raise last_err


def _rewrite_last_answer_for_lang(
    gclient,
    GEMINI_MODEL,
    GEN_CFG,
    last_answer: str,
    target_lang: str = "en",
) -> str:
    last_answer = (last_answer or "").strip()
    if not last_answer:
        return ""

    target_lang = (target_lang or "en").lower()

    if target_lang.startswith("en"):
        instr = (
            "Rewrite the following answer in clear, natural English. "
            "Keep the meaning and important details, but you can shorten slightly if needed.\n"
            "Answer ONLY in English.\n"
        )
    else:
        instr = (
            "Viết lại câu trả lời sau bằng tiếng Việt rõ ràng, tự nhiên. "
            "Giữ nguyên ý chính và các chi tiết quan trọng, có thể rút gọn nhẹ nếu cần.\n"
            "Chỉ trả lời bằng tiếng Việt.\n"
        )

    prompt = f"""{instr}
    [ANSWER]
    {last_answer}
    """

    contents = [
        {
            "role": "user",
            "parts": [{"text": prompt}],
        }
    ]

    out, _ = _safe_gemini_generate(gclient, GEMINI_MODEL, contents, GEN_CFG)
    return (out or "").strip()


def _analyze_followup_and_lang(
    gclient,
    GEMINI_MODEL,
    GEN_CFG,
    user_text: str,
    history_msgs: List[Dict[str, str]],
    ctx: Dict[str, Any],
):
    user_text = (user_text or "").strip()
    if not user_text:
        return {"mode": "pass", "output": "", "set_lang": None}

    last_ans = _get_last_assistant_message(history_msgs)
    if not last_ans:
        return {"mode": "pass", "output": "", "set_lang": None}

    current_lang = (ctx or {}).get("lang") or "vi"
    lower_ut = user_text.lower()

    vi_to_en_patterns = [
        "viết bằng tiếng anh",
        "viết lại bằng tiếng anh",
        "viết tiếng anh",
        "dịch sang tiếng anh",
        "dịch đoạn trên sang tiếng anh",
        "dịch nội dung trên sang tiếng anh",
        "rewrite in english",
        "write in english",
        "answer in english",
        "translate to english",
    ]
    en_to_vi_patterns = [
        "viết bằng tiếng việt",
        "viết lại bằng tiếng việt",
        "dịch sang tiếng việt",
        "dịch đoạn trên sang tiếng việt",
        "dịch nội dung trên sang tiếng việt",
        "rewrite in vietnamese",
        "translate to vietnamese",
    ]

    if any(p in lower_ut for p in vi_to_en_patterns):
        out_text = _rewrite_last_answer_for_lang(
            gclient, GEMINI_MODEL, GEN_CFG, last_ans, target_lang="en"
        )
        return {
            "mode": "edit",
            "output": out_text,
            "set_lang": "en",
        }

    if any(p in lower_ut for p in en_to_vi_patterns):
        out_text = _rewrite_last_answer_for_lang(
            gclient, GEMINI_MODEL, GEN_CFG, last_ans, target_lang="vi"
        )
        return {
            "mode": "edit",
            "output": out_text,
            "set_lang": "vi",
        }

    prompt = f"""{sys_instr_warehouse}

    [CURRENT_PREFERRED_LANG]
    {current_lang}

    [PREVIOUS_ANSWER]
    {last_ans}

    [USER_REQUEST]
    {user_text}
    """

    contents = [
        {
            "role": "user",
            "parts": [{"text": prompt}],
        }
    ]

    raw, _ = _safe_gemini_generate(gclient, GEMINI_MODEL, contents, GEN_CFG)
    raw = (raw or "").strip()

    try:
        data = json.loads(raw)
    except Exception:
        data = {}

    if not isinstance(data, dict):
        data = {}

    mode = (data.get("mode") or "pass").lower()
    output = (data.get("output") or "").strip()
    set_lang = data.get("set_lang", None)

    if set_lang is not None:
        if isinstance(set_lang, str):
            set_lang = set_lang.lower()
            if set_lang not in ("vi", "en"):
                set_lang = None
        else:
            set_lang = None

    if set_lang is None:
        detected = _auto_detect_lang(user_text)
        if detected in ("vi", "en") and detected != current_lang:
            set_lang = detected

    if mode != "edit" or not output:
        return {"mode": "pass", "output": "", "set_lang": set_lang}

    return {"mode": "edit", "output": output, "set_lang": set_lang}


# ======================================================================
# 4. FILE UPLOAD (VS session) + STRIP SOURCE
# ======================================================================
def strip_source_citations(text: str) -> str:
    if not text:
        return text

    text = re.sub(
        r'\bsource["\']?\s*:\s*["\']?[^"\n]*\.(?:txt|docx|pdf)["\']?',
        '',
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r'[\(\[\{]\s*(?:source|chunk_id|metadata|norm_case)[\s:][^\)\]\}]{0,400}[\)\]\}]',
        '',
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r'\(\s*[Ll]ink\s*:\s*(?!https?://)[^)]*\.(?:txt|docx|pdf)\s*\)',
        '',
        text,
        flags=re.DOTALL,
    )

    text = re.sub(
        r'\b[Ll]ink\s*:\s*(?!https?://)[^\n]*\.(?:txt|docx|pdf)',
        '',
        text,
        flags=re.DOTALL,
    )

    text = re.sub(
        r'\b[Nn]gu[oơ]n\s*:\s*[^\n]*\.(?:txt|docx|pdf)',
        '',
        text,
        flags=re.DOTALL,
    )

    text = re.sub(r'[ \t]{2,}', ' ', text)
    text = re.sub(r'\s+([.,!?;:])', r'\1', text)
    text = re.sub(r'\(\s*\)', '', text)
    text = re.sub(r'\.{3,}', '...', text)
    text = text.strip()

    return text


def _extract_text_from_uploaded_file(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext == ".txt":
        try:
            with open(path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception:
            return ""

    if ext == ".docx":
        try:
            d = docx_lib.Document(path)
            lines = []
            for p in d.paragraphs:
                t = (p.text or "").strip()
                if t:
                    lines.append(t)
            for tbl in d.tables:
                for row in tbl.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    line = " | ".join(c for c in cells if c)
                    if line:
                        lines.append(line)
            return "\n".join(lines)
        except Exception:
            return ""

    if ext == ".pdf":
        try:
            doc = fitz.open(path)
            pages = []
            for p in doc:
                t = p.get_text()
                if t.strip():
                    pages.append(t)
            doc.close()
            return "\n".join(pages)
        except Exception:
            return ""

    return ""


def _split_text_to_chunks(text: str, max_chars: int = 800, overlap: int = 200) -> List[str]:
    text = (text or "").strip()
    if not text:
        return []

    chunks: List[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + max_chars, n)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == n:
            break
        start = max(0, end - overlap)
    return chunks


def _build_session_doc_vs(
    session_id: str,
    embeddings,
    file_paths: List[str],
) -> Optional[FAISS]:
    if not session_id or not file_paths:
        return None

    texts: List[str] = []
    metas: List[Dict[str, Any]] = []

    for path in file_paths:
        if not path or not os.path.exists(path):
            continue
        try:
            raw = _extract_text_from_uploaded_file(path)
        except Exception:
            raw = ""
        if not raw.strip():
            continue

        chunks = _split_text_to_chunks(raw, max_chars=800, overlap=200)
        src = os.path.basename(path)
        for ch in chunks:
            texts.append(ch)
            metas.append({"source": src, "session_id": session_id})

    if not texts:
        return None

    vs = FAISS.from_texts(texts, embeddings, metadatas=metas)
    SESSION_DOC_VS[session_id] = vs
    return vs


def _get_session_doc_vs(
    session_id: str,
    embeddings,
    file_paths: List[str],
) -> Optional[FAISS]:
    if not session_id:
        return None

    vs = SESSION_DOC_VS.get(session_id)
    if vs is not None:
        return vs

    if not file_paths:
        return None

    return _build_session_doc_vs(session_id, embeddings, file_paths)


def _clear_session_doc_vs(session_id: str) -> None:
    SESSION_DOC_VS.pop(session_id, None)


def _extract_src_and_url(meta: dict) -> tuple[str, str]:
    if not isinstance(meta, dict):
        return "", ""

    src_name = (
        meta.get("title")
        or meta.get("source")
        or meta.get("filename")
        or ""
    )

    if not src_name and meta.get("path"):
        src_name = os.path.basename(meta["path"])

    url = meta.get("url") or meta.get("link") or meta.get("path") or ""

    return src_name or "", url or ""


# ======================================================================
# 5. GEMINI HISTORY + RAG + DB
# ======================================================================
def _to_gemini_history_no_system(history_msgs: List[Dict[str, str]]) -> List[Dict]:
    out = []
    for m in history_msgs:
        role = m.get("role", "user")
        content = (m.get("content") or "").strip()
        if not content:
            continue
        role = "model" if role == "assistant" else "user"
        out.append({"role": role, "parts": [{"text": content}]})
    return out


def _rewrite_db_answer(
    gclient,
    GEMINI_MODEL,
    GEN_CFG,
    answer_text: str,
    lang_hint: str | None = None,
) -> str:
    if not REWRITE_DB_WITH_LLM:
        return answer_text

    lang_hint = (lang_hint or "vi").lower()
    if lang_hint.startswith("en"):
        lang_rule = "\nTrả lời bằng tiếng Anh. Giữ nguyên số liệu, mã đơn, phần trăm, mã lô."
    else:
        lang_rule = (
            "\nTrả lời bằng tiếng Việt. Giữ nguyên số liệu, mã đơn, phần trăm, mã lô."
        )

    guard = (
        SYSTEM_PRIMER_WAREHOUSE
        + lang_rule
        + "\nYÊU CẦU: giữ NGUYÊN số liệu, số tiền, phần trăm, mã đơn, mã lô.\n"
        "Không thêm bớt số. Nếu một số không cần thiết có quyền bỏ. "
        "Chỉ chỉnh lại câu chữ cho tự nhiên, gọn trong 1-2 đoạn & giữ nguyên bảng ASCII nếu có."
    )
    contents = [
        {
            "role": "user",
            "parts": [
                {
                    "text": f"[INSTRUCTION]\n{guard}\n\n[TEXT]\n{answer_text}",
                }
            ],
        }
    ]
    try:
        out, _ = _safe_gemini_generate(gclient, GEMINI_MODEL, contents, GEN_CFG)
        out = strip_source_citations(out or "")
        has_num_old = re.search(r"\d", answer_text or "") is not None
        has_num_new = re.search(r"\d", out or "") is not None
        if has_num_old and not has_num_new:
            return answer_text
        return out or answer_text
    except Exception:
        return answer_text


def _rag_answer(
    user_text: str,
    gclient,
    GEMINI_MODEL,
    GEN_CFG,
    embeddings,
    retriever,
    lang_hint: str | None = None,
    session_id: str | None = None,
    session_doc_paths: Optional[List[str]] = None,
) -> tuple[str, dict]:
    session_doc_paths = session_doc_paths or []
    start_total = time.time()

    session_docs = []
    search_time_session = 0.0

    if session_doc_paths:
        try:
            vs = _get_session_doc_vs(session_id, embeddings, session_doc_paths)
            if vs is not None:
                t0 = time.time()
                session_docs = vs.similarity_search(user_text, k=15)
                search_time_session = time.time() - t0

                if len(session_docs) < 3:
                    session_docs = []
        except Exception as e:
            print(f"[RAG] Lỗi khi tìm trong session docs: {e}")
            session_docs = []

    if not session_docs:
        try:
            t0 = time.time()
            global_docs = retriever.get_relevant_documents("query: " + user_text)
            search_time_global = time.time() - t0
        except Exception:
            global_docs = []
        all_candidates = global_docs
        search_time_session = 0.0
    else:
        all_candidates = session_docs
        search_time_global = 0.0

    rerank_time = 0.0
    docs_text = ""
    try:
        t0 = time.time()
        ranked = rerank(user_text, all_candidates, top_k=TOP_K)
        parts = []
        for d, _score in ranked:
            txt = d.page_content
            if txt.lower().startswith("passage: "):
                txt = txt[len("passage: "):]
            src_name, url = _extract_src_and_url(d.metadata)
            suffix = f" (Nguồn: {src_name})" if src_name else ""
            if url and "http" in url:
                suffix += f" | Link: {url}"
            parts.append(txt + suffix)
        docs_text = "\n\n---\n\n".join(parts) if parts else ""
        rerank_time = time.time() - t0
    except Exception as e:
        print(f"[RAG] Lỗi rerank: {e}")

    if docs_text:
        context_hint = f"Dữ liệu tham khảo:\n{docs_text}"
    else:
        context_hint = "(Không tìm thấy thông tin phù hợp từ tài liệu.)"

    lang = (lang_hint or "vi").lower()
    if lang.startswith("en"):
        lang_instruction = "\nTrả lời bằng tiếng Anh, rõ ràng, chuyên nghiệp."
    else:
        lang_instruction = "\nTrả lời bằng tiếng Việt, tự nhiên, dễ hiểu."

    system_prompt = f"""{SYSTEM_PRIMER_WAREHOUSE}
{lang_instruction}

{context_hint}
"""

    hist_msgs = [
        {"role": m["role"], "content": m["content"]}
        for m in get_history()
        if m["role"] in ("user", "assistant")
    ]
    contents = _to_gemini_history_no_system(hist_msgs)
    contents.append({
        "role": "user",
        "parts": [{"text": f"[SYSTEM]\n{system_prompt}\n\n[USER]\n{user_text}"}]
    })

    try:
        result, llm_time = _safe_gemini_generate(gclient, GEMINI_MODEL, contents, GEN_CFG)
    except Exception as e:
        error_msg = "Hệ thống đang bận. Vui lòng thử lại sau."
        if "quota" in str(e).lower() or "resource_exhausted" in str(e).lower():
            error_msg = "Hệ thống LLM tạm thời hết quota. Vui lòng thử lại sau ít phút."
        result, llm_time = error_msg, 0.0

    final_answer = strip_source_citations(result.strip())

    timing = {
        "embedding": 0.0,
        "search": round(search_time_global + search_time_session + rerank_time, 2),
        "llm": round(llm_time, 2),
        "total": round(time.time() - start_total, 2),
    }

    return final_answer, timing


# ======================================================================
# 6. HÀM CHÍNH: xử lý 1 lượt chat (sync)
# ======================================================================
def handle_chat_request(
    user_text: str,
    session_id: str,
    principal: Any,
    bm25_folder: Optional[Path] = None,
    corpus_path: Optional[Path] = None,
) -> dict:
    """
    Trả về dict để API dùng jsonify.
    bm25_folder / corpus_path:
      - mặc định: dùng DATA_DIR + FAISS_ALL_DIR/corpus.jsonl
      - có thể override cho HR/Accountant/chat khác.
    """
    user_text = (user_text or "").strip()
    if not user_text:
        return {"ok": False, "error": "Missing message"}

    # ====== Chuẩn bị session_id (dùng Flask session) ======
    if not session_id:
        try:
            info = create_new_session(
                session.get("user") or "", title="Cuộc trò chuyện mới"
            )
            session_id = info.get("session_id") or ""
        except Exception:
            session_id = ""

    # Ghi history user
    add_message("user", user_text, session_id=session_id)

    full_start = time.time()

    bm25_folder = bm25_folder or DATA_DIR_WAREHOUSE
    corpus_path = corpus_path or (FAISS_ALL_DIR_WAREHOUSE / "corpus.jsonl")

    embeddings, vector_store, retriever, gclient, GEMINI_MODEL, GEN_CFG = LLM_model(
        bm25_folder=bm25_folder,
        corpus_path=corpus_path,
    )

    # ==== LOAD CTX & phân tích ngữ cảnh + ngôn ngữ ====
    ctx = _load_ctx(session_id)
    history_msgs = _get_history_msgs_for_ctx()

    analysis = _analyze_followup_and_lang(
        gclient, GEMINI_MODEL, GEN_CFG, user_text, history_msgs, ctx
    )

    new_lang = analysis.get("set_lang")
    if new_lang in ("vi", "en"):
        ctx["lang"] = new_lang
    else:
        if "lang" not in ctx:
            ctx["lang"] = _auto_detect_lang(user_text)

    # ==== Editor: chỉ chỉnh sửa/dịch câu trả lời trước ====
    if analysis.get("mode") == "edit":
        final_answer = strip_source_citations(analysis.get("output") or "")
        add_message("assistant", final_answer, session_id=session_id)

        try:
            ctx_to_save = dict(ctx)
            ctx_to_save["last_branch"] = "editor"
            update_ctx(session_id, ctx_to_save)
        except Exception:
            pass

        elapsed = time.time() - full_start
        return {
            "ok": True,
            "answer": final_answer,
            "session_id": session_id,
            "branch": "editor",
            "meta": {"editor": {"mode": "edit", "set_lang": ctx.get("lang")}},
            "timing": {
                "total": round(elapsed, 2),
                "embedding": 0.0,
                "search": 0.0,
                "llm": round(elapsed, 2),
            },
        }

    # ==== XÓA FILE UPLOAD: logic vẫn ở đây, API chỉ gọi ====
    normalized = user_text.lower().strip()
    if normalized in ["xóa file", "xoá file", "hủy file", "huy file", "delete file"]:
        session.pop("uploaded_files", None)
        session.pop("session_docs_dirty", None)
        _clear_session_doc_vs(session_id)

        final_answer = (
            "Đã xoá danh sách file đã upload trong phiên hiện tại. "
            "Các câu hỏi tiếp theo sẽ chỉ dùng DB/RAG global."
        )
        add_message("assistant", final_answer, session_id=session_id)
        elapsed = time.time() - full_start
        return {
            "ok": True,
            "answer": final_answer,
            "session_id": session_id,
            "branch": "doc_clear",
            "meta": {"doc_qa": {"cleared": True}},
            "timing": {
                "total": round(elapsed, 2),
                "embedding": 0.0,
                "search": 0.0,
                "llm": 0.0,
            },
        }

    # ==== LẤY DANH SÁCH FILE UPLOAD TỪ session ====
    uploaded_files_meta = session.get("uploaded_files", [])
    session_doc_paths: list[str] = []
    if isinstance(uploaded_files_meta, list):
        for item in uploaded_files_meta:
            if isinstance(item, dict):
                p = item.get("path")
                if p and isinstance(p, str) and os.path.exists(p):
                    session_doc_paths.append(p)

    # ==== NHÁNH DB TRƯỚC ====
    handled, db_answer, meta = handle_db_message(user_text, session_id, principal=principal)
    branch = "db" if handled else "rag"
    timing = {"embedding": 0.0, "search": 0.0, "llm": 0.0}

    if handled:
        final_answer = _rewrite_db_answer(
            gclient, GEMINI_MODEL, GEN_CFG, db_answer, lang_hint=ctx.get("lang")
        )
        try:
            ctx_to_save = dict(ctx)
            if isinstance(meta, dict):
                if meta.get("intent"):
                    ctx_to_save["last_intent"] = meta.get("intent")
                if meta.get("time_window"):
                    ctx_to_save["time_window"] = meta.get("time_window")
            ctx_to_save["last_branch"] = "db"
            update_ctx(session_id, ctx_to_save)
        except Exception:
            pass
    else:
        # ==== NHÁNH RAG ====
        final_answer, timing = _rag_answer(
            user_text,
            gclient,
            GEMINI_MODEL,
            GEN_CFG,
            embeddings,
            retriever,
            lang_hint=ctx.get("lang"),
            session_id=session_id,
            session_doc_paths=session_doc_paths,
        )
        try:
            ctx_to_save = dict(ctx)
            ctx_to_save["last_branch"] = "rag"
            if session_doc_paths:
                ctx_to_save["last_docs"] = [os.path.basename(p) for p in session_doc_paths]
            update_ctx(session_id, ctx_to_save)
        except Exception:
            pass

    add_message("assistant", final_answer, session_id=session_id)

    # ==== SMART SUMMARY (giữ nguyên logic cũ, chỉ move vào đây) ====
    try:
        history = get_history()
        user_messages_count = len([m for m in history if m.get("role") == "user"])
        if user_messages_count >= 7 and user_messages_count % 7 == 0:
            summary_prompt = (
                "Bạn là trợ lý nội bộ cực kỳ thông minh. Hãy tóm tắt cuộc trò chuyện sau đây sao cho "
                "vẫn giữ được toàn bộ ngữ cảnh quan trọng, chi tiết quy trình, biểu mẫu, số tiền, "
                "tên người/tên phòng ban đã nhắc đến (nếu có). Viết dưới góc nhìn thứ nhất như người dùng đang nói.\n"
                "Độ dài: 3–5 câu, tối đa 250 từ. Bắt đầu bằng 'Người dùng đang hỏi về...'\n\n"
                "Cuộc trò chuyện gần nhất:\n"
                + "\n".join([f"{m['role']}: {m['content'][:800]}" for m in history[-15:]])
                + "\n\nTóm tắt:"
            )

            from google.generativeai import GenerativeModel

            summary_model = GenerativeModel("gemini-2.0-flash")
            summary_resp = summary_model.generate_content(
                summary_prompt,
                generation_config={
                    "temperature": 0.4,
                    "max_output_tokens": 300,
                    "top_p": 0.95,
                },
            )
            summary_text = summary_resp.text.strip()

            redis_client.delete(f"history:{session_id}")
            redis_client.delete(f"ctx:{session_id}")
            redis_client.delete(f"session_doc_vs:{session_id}")
            add_message("system", f"[TÓM TẮT NGỮ CẢNH]: {summary_text}", session_id=session_id)
            add_message("user", user_text, session_id=session_id)

            print(f"[SMART SUMMARY] Đã tóm tắt & reset Redis cho session {session_id}")

    except Exception as e:
        print(f"[SUMMARY ERROR] {e}")

    elapsed = time.time() - full_start
    return {
        "ok": True,
        "answer": final_answer,
        "session_id": session_id,
        "branch": branch,
        "meta": meta,
        "timing": {
            "total": round(elapsed, 2),
            **timing,
        },
    }